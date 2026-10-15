"""
L'export Excel del GANTT (`GET /api/gantt/export-excel`) contro la pagina.

Gira su gantt_test (fixture `db_test`). Il criterio è uno solo: l'Excel
contiene i task che la pagina GANTT mostra con gli stessi filtri. La pagina
legge /api/gantt/strutturato con `stato` e `progetto_id` e filtra il profilo
nel browser, su `profilo_richiesto` (pages/Gantt.jsx): qui si rifà lo stesso
conto sul payload di /strutturato e lo si confronta con le righe dell'Excel.
"""

import io
from collections import Counter

import pytest
from openpyxl import load_workbook

from routes.gantt import export_gantt_excel, gantt_strutturato
from stati_colori import carica_stati

pytestmark = pytest.mark.usefixtures("db_test")


def _task_della_pagina(stato, progetto_id, profilo):
    progetti = gantt_strutturato(stato=stato, progetto_id=progetto_id, _=None)
    return Counter(
        (p["nome"], t["nome"])
        for p in progetti for f in p["fasi"] for t in f["tasks"]
        if not profilo or t["profilo_richiesto"] == profilo
    )


def _excel(stato, progetto_id, profilo):
    risposta = export_gantt_excel(stato=stato, progetto_id=progetto_id,
                                  profilo=profilo, _=None)
    return risposta, load_workbook(io.BytesIO(risposta.body))


def _task_dell_excel(wb):
    righe = list(wb["Dati GANTT"].iter_rows(min_row=2, values_only=True))
    return Counter((r[0], r[1]) for r in righe)


def _un_progetto_e_un_profilo():
    progetti = gantt_strutturato(stato="attivi", progetto_id=None, _=None)
    profili = Counter(t["profilo_richiesto"] for p in progetti for f in p["fasi"]
                      for t in f["tasks"] if t["profilo_richiesto"])
    return progetti[0]["id"], profili.most_common(1)[0][0]


@pytest.mark.parametrize("stato", ["attivi", "in esecuzione", "sospeso"])
def test_stessi_task_della_pagina_per_ogni_stato(stato):
    _, wb = _excel(stato, None, None)
    attesi = _task_della_pagina(stato, None, None)
    assert attesi, f"il seed non ha task per lo stato «{stato}»"
    assert _task_dell_excel(wb) == attesi


def test_niente_filtro_proprio_sui_sospesi():
    """Prima l'Excel del portafoglio toglieva i progetti Sospesi anche quando
    la pagina li mostrava: con «Solo Sospeso» il file usciva vuoto."""
    _, wb = _excel("sospeso", None, None)
    assert sum(_task_dell_excel(wb).values()) > 0


def test_progetto_e_profilo_come_la_pagina():
    progetto_id, profilo = _un_progetto_e_un_profilo()
    for argomenti in [("attivi", progetto_id, None), ("attivi", None, profilo),
                      ("attivi", progetto_id, profilo)]:
        _, wb = _excel(*argomenti)
        assert _task_dell_excel(wb) == _task_della_pagina(*argomenti), argomenti


def test_filtri_che_non_lasciano_task_danno_un_file_vuoto_non_un_errore():
    risposta, wb = _excel("attivi", None, "Profilo che non esiste")
    assert risposta.status_code == 200
    assert _task_dell_excel(wb) == Counter()


def test_nome_del_file_come_pdf_e_png():
    progetto_id, _ = _un_progetto_e_un_profilo()
    risposta, _ = _excel("attivi", None, None)
    assert "filename=gantt_portafoglio_" in risposta.headers["content-disposition"]
    risposta, _ = _excel("attivi", progetto_id, None)
    assert f"filename=gantt_{progetto_id}_" in risposta.headers["content-disposition"]


def test_colori_degli_stati_dal_file_condiviso():
    stati = carica_stati()
    _, wb = _excel("attivi", None, None)
    ws = wb["Dati GANTT"]
    colonna_stato = [c.value for c in ws[1]].index("Stato") + 1
    for riga in range(2, ws.max_row + 1):
        cella = ws.cell(row=riga, column=colonna_stato)
        atteso = stati["task"].get(cella.value, stati["ignoto"])["colore"].lstrip("#")
        assert cella.fill.start_color.rgb[-6:].lower() == atteso.lower(), cella.value
